"""Build the written analysis from the executed R output tables."""
from pathlib import Path
import argparse
import base64
import csv
import html
import shutil
import statistics
import subprocess
import tempfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Florien Siakoua Toukam"
TITLE = "Predicting Financial Distress in Polish Companies"


def read(name):
    with (ROOT / "outputs/tables" / (name + ".csv")).open() as f:
        return list(csv.DictReader(f))


def number(row, key):
    return float(row[key])


def pct(value, digits=1):
    return f"{float(value) * 100:.{digits}f}%"


def num(value, digits=3):
    return f"{float(value):.{digits}f}"


def integer(value):
    return f"{int(float(value)):,}"


def paragraph(text):
    return ("p", text)


def table(headers, rows):
    return ("table", headers, rows)


def picture(name, caption):
    return ("figure", name, caption)


perf = read("horizon-holdout-performance")
primary = [x for x in perf if x["horizon_years"] == "1"]
lr, rf = primary
bands = read("holdout-risk-bands")
edges = read("development-risk-band-boundaries")
selection = [x for x in read("complete-case-selection-diagnostic") if x["file"] == "5year.arff"]
model_selection = [x for x in selection if x["feature_set"] == "Model 16 ratios"]
collinearity = read("development-collinearity-diagnostics")
inventory = read("data-inventory")
drivers = read("logistic-risk-drivers")
cv = read("cross-validation-fold-results")
ci = read("holdout-metric-confidence-intervals")
delta = read("paired-holdout-auc-difference")[0]
lift = read("holdout-lift-and-capture")
cal = read("holdout-calibration-scores")
sensitivity = read("financial-ratio-sensitivity")
splits = [x for x in read("development-holdout-samples") if x["file"] == "5year.arff"]
cv_means = {m: statistics.mean(float(x["roc_auc"]) for x in cv if x["horizon_years"] == "1" and x["model"] == m) for m in (lr["model"], rf["model"])}
lr_top10 = lift[0]
selection_text = (
    f"Requiring complete values in the 16 model ratios would exclude {integer(model_selection[1]['removed_n'])} of "
    f"{integer(model_selection[1]['n'])} bankrupt development records ({pct(model_selection[1]['removed_fraction'])}), "
    f"compared with {integer(model_selection[0]['removed_n'])} of {integer(model_selection[0]['n'])} non-bankrupt records "
    f"({pct(model_selection[0]['removed_fraction'])}). Training-only median imputation and missing indicators retain every record."
)

headline = (
    f"The one-year holdout supports both models for prioritizing financial review. "
    f"Ridge logistic regression has ROC AUC {num(lr['roc_auc'])}; the balanced random forest has {num(rf['roc_auc'])}. "
    f"The forest's higher AUC does not establish a clear overall advantage: the paired difference's 95% interval includes zero, "
    f"while logistic regression has higher average precision and lower probability error. "
    f"I would retain logistic regression as the interpretable reference model and use the forest as a challenger."
)
policy_text = (
    f"An assumed capacity to review 20% of development observations produces a {pct(lr['threshold'], 2)} logistic screening threshold "
    f"and a {pct(rf['threshold'], 2)} forest threshold. Applied unchanged to {integer(lr['n'])} holdout records, "
    f"logistic regression flags {integer(number(lr, 'tp') + number(lr, 'fp'))} records and identifies {integer(lr['tp'])} of "
    f"{integer(lr['bankruptcies'])} bankruptcies. The forest flags {integer(number(rf, 'tp') + number(rf, 'fp'))} and identifies "
    f"{integer(rf['tp'])}. These are initial review thresholds, not automatic credit decisions."
)
band_text = (
    f"Development-defined logistic bands produce observed holdout bankruptcy rates of {pct(bands[0]['observed_bankruptcy_rate'])}, "
    f"{pct(bands[1]['observed_bankruptcy_rate'])} and {pct(bands[2]['observed_bankruptcy_rate'])} for Low, Moderate and High risk. "
    f"The High band contains {integer(bands[2]['n'])} records and {integer(lr['tp'])} observed bankruptcies."
)
comparison = (
    f"At the development-selected thresholds, the forest captures one additional bankruptcy with 20 fewer false positives. "
    f"Across all score thresholds, logistic regression has higher average precision "
    f"({num(lr['average_precision'])} versus {num(rf['average_precision'])}). "
    f"Its top decile captures {pct(lr_top10['capture_rate'])} of holdout bankruptcies. "
    f"The choice therefore depends on review capacity, the quality of risk ranking at the top of the list and interpretability."
)
limitations = (
    "These are historical, selected Polish company records, not a current credit book. Firm identifiers, observation dates, "
    "sector classifications, exposures and recovery outcomes are absent. Exact duplicate grouping reduces a known leakage risk, "
    "but cannot eliminate unidentified repeated-company dependence. The holdout is random rather than calendar-time based. "
    "Calibration is relative to the dataset's event mix, not a validated current-market probability of default."
)
driver_text = (
    "Higher sales growth, working capital relative to assets, company size, and profitability are associated with lower estimated "
    "bankruptcy odds in the regularized model. A higher liabilities-to-assets ratio is associated with higher odds. "
    "These directions are stable across the ten outer fits for the six ratios shown below. "
    "Coefficients are conditional on the other inputs; they do not identify causal effects."
)
metric_rows = []
for label, key, fmt in [
    ("ROC AUC", "roc_auc", num), ("Average precision", "average_precision", num),
    ("Bankruptcy recall", "recall", pct), ("Specificity", "specificity", pct),
    ("Precision", "precision", pct), ("Accuracy", "accuracy", pct),
    ("Balanced accuracy", "balanced_accuracy", pct), ("F1", "f1", num),
    ("Brier score", "brier", lambda x: num(x, 4)), ("Review share", "review_rate", pct),
    ("True positives", "tp", integer), ("False positives", "fp", integer),
    ("False negatives", "fn", integer), ("True negatives", "tn", integer),
]:
    metric_rows.append([label, fmt(lr[key]), fmt(rf[key])])

sections = [
    ("Credit risk assessment", [
        paragraph(headline), paragraph(policy_text),
        table(["Primary holdout", "Ridge logistic", "Balanced forest"], [
            ["ROC AUC", num(lr['roc_auc']), num(rf['roc_auc'])],
            ["Average precision", num(lr['average_precision']), num(rf['average_precision'])],
            ["Bankruptcy recall", pct(lr['recall']), pct(rf['recall'])],
            ["Precision", pct(lr['precision']), pct(rf['precision'])],
        ]),
        paragraph(band_text),
        paragraph("The analysis examines data quality and financial ratios, then evaluates two models under a shared development-and-holdout design. "
                  "The objective is to prioritize records for financial investigation while making the review workload and missed cases explicit."),
    ]),
    ("Data and forecasting horizons", [
        paragraph("The data come from Sebastian Tomczak's Polish Companies Bankruptcy dataset, published by the UCI Machine Learning Repository. "
                  "Each record contains 64 financial predictors and a bankruptcy indicator. All five supplied ARFF files are retained unchanged; "
                  "their SHA-256 checksums identify the exact inputs."),
        table(["File", "Horizon", "Records", "Bankruptcies", "Event rate"], [
            [x['file'], x['horizon_years'] + (" year" if x['horizon_years'] == '1' else " years"), integer(x['observations']), integer(x['bankruptcies']), pct(x['bankruptcy_rate'])] for x in inventory
        ]),
        paragraph("The filename is not the remaining prediction horizon: 1year.arff refers to the first forecasting-period year and labels bankruptcy after five years. "
                  "The primary one-year-ahead sample is 5year.arff. Each file is modeled separately; the analysis does not pool them or assume they contain distinct firms."),
        paragraph(f"The primary split contains {integer(splits[0]['n'])} development records with {integer(splits[0]['bankruptcies'])} bankruptcies "
                  f"and {integer(splits[1]['n'])} holdout records with {integer(splits[1]['bankruptcies'])} bankruptcies. "
                  f"Its {integer(inventory[4]['duplicate_excess'])} excess exact predictor matches are retained but assigned to the same split as their matching records."),
        paragraph("Missing values are handled inside each training sample. Median imputation and missing-value indicators preserve incomplete observations; "
                  "training-percentile clipping limits the influence of extreme ratio values. The raw files are never rewritten."),
    ]),
    ("Data quality and sample inclusion", [
        paragraph(selection_text),
        picture("complete-case-selection-impact", "One-year development records excluded by complete-case requirements. No records are excluded from model fitting."),
        paragraph(f"Bankruptcies account for {pct(inventory[4]['bankruptcy_rate'])} of the one-year sample. "
                  f"Labeling every record non-bankrupt would achieve {pct(1 - number(inventory[4], 'bankruptcy_rate'))} accuracy but zero bankruptcy recall. "
                  "Evaluation therefore includes precision-recall performance, screening capture and false-positive review volume alongside ROC AUC."),
        paragraph(f"After development-only imputation and clipping, variance inflation factors across the 16 ratios range from "
                  f"{num(min(number(x, 'vif') for x in collinearity), 2)} to {num(max(number(x, 'vif') for x in collinearity), 2)}. "
                  "The development correlation matrix and VIF table document predictor dependence. Ridge regularization stabilizes fitting; "
                  "the coefficients still describe associations conditional on the other ratios, not independent causal effects."),
    ]),
    ("Model development and validation", [
        paragraph("The comparison uses the same 16 finance-defined ratios for both models. Inputs cover profitability, leverage, liquidity, "
                  "capital accumulation, sales growth, size and working-capital efficiency. The set is fixed before evaluation, without outcome-based feature selection."),
        paragraph("Exact matches across all 64 raw predictors form groups. Approximately 20% of groups within each outcome stratum are reserved for holdout evaluation. "
                  "Development uses two repeats of grouped five-fold cross-validation. No matching predictor group crosses a training/validation or development/holdout boundary."),
        paragraph("Every fit estimates medians and 1st/99th percentile clipping limits on its own training observations. Each of the 16 ratios also receives a missing indicator. "
                  "Ridge logistic regression standardizes inputs and uses a fixed penalty of 0.01. The random forest uses 500 trees, five candidate variables at each split, "
                  "minimum terminal-node size five, and equal bootstrap counts from each outcome class."),
        paragraph("Each outer training fold runs a separate three-fold procedure to obtain calibration scores. Platt calibration is fitted to those inner out-of-fold scores "
                  "at the natural event prevalence, then applied to the outer validation predictions. This keeps calibration separate from validation outcomes. "
                  "The two outer validation scores per development record are averaged for screening-policy and risk-band design."),
        paragraph("Final calibrators use development out-of-fold scores. Final base models use all development records. Their fitted transformations, calibrators, "
                  "screening thresholds and risk-band boundaries are fixed before evaluating the holdout. The forest's balanced vote fraction is not treated as an unadjusted probability."),
        paragraph("Model specifications are fixed rather than chosen by a large tuning search. Base seed 5442026, the file-specific seed schedule and package versions are recorded. "
                  "The 95% holdout intervals use 500 paired, outcome-stratified bootstrap samples of predictor groups. They are conditional on the fitted models and the chosen split."),
    ]),
    ("One year holdout performance", [
        table(["Metric", "Ridge logistic", "Balanced forest"], metric_rows),
        paragraph(comparison),
        paragraph(f"The forest-minus-logistic AUC difference is {num(delta['estimate'])}, with a 95% interval from {num(delta['lower_95'])} to {num(delta['upper_95'])}. "
                  f"Mean outer-fold AUCs are {num(cv_means[lr['model']])} for logistic regression and {num(cv_means[rf['model']])} for the forest. "
                  "Cross-validation supports the forest's modest discrimination advantage, but does not establish superiority on every metric."),
    ]),
    ("Risk ranking and calibration", [
        picture("holdout-roc-and-precision-recall", "One-year holdout curves. Average precision uses the step precision-recall definition, not trapezoidal PR AUC."),
        paragraph(f"The holdout bankruptcy prevalence is {pct(number(lr, 'bankruptcies') / number(lr, 'n'))}. "
                  f"The highest-scoring logistic decile contains {integer(lr_top10['selected_n'])} records, captures {pct(lr_top10['capture_rate'])} of bankruptcies, "
                  f"and achieves {num(lr_top10['lift'], 2)} times the full-sample event rate. "
                  "Top-decile and top-quintile diagnostics use holdout rank, whereas the screening policy uses fixed development thresholds."),
        table(["Probability error", "Ridge logistic", "Balanced forest"], [
            ["Raw Brier score", num(cal[0]['raw_brier'], 4), num(cal[1]['raw_brier'], 4)],
            ["Calibrated Brier score", num(cal[0]['calibrated_brier'], 4), num(cal[1]['calibrated_brier'], 4)],
            ["Development-prevalence baseline", num(cal[0]['development_prevalence_baseline_brier'], 4), num(cal[1]['development_prevalence_baseline_brier'], 4)],
        ]),
        paragraph("Calibration substantially reduces the forest's probability error. It slightly increases logistic Brier error on this holdout, "
                  "so it is not described as a universal improvement. Calibration-bin plots are included with the supporting figures. "
                  "A lower Brier score is better; the constant baseline uses development prevalence rather than learning the holdout event rate."),
    ]),
    ("Screening policy and risk segmentation", [
        paragraph(policy_text),
        table(["Logistic band", "Records", "Average score", "Observed rate"], [
            [x['band'], integer(x['n']), pct(x['average_predicted_risk']), pct(x['observed_bankruptcy_rate'])] for x in bands
        ]),
        paragraph(f"The Low-to-Moderate boundary is {pct(edges[0]['threshold'], 2)} and the Moderate-to-High boundary is {pct(edges[1]['threshold'], 2)}. "
                  "They correspond approximately to the median and 80th percentile of cross-fitted development scores. "
                  "High is a relative screening band, not a regulatory grade or an independently validated rating."),
        picture("development-screening-capacity", "Development policy sensitivity. The vertical line marks the assumed 20% review budget, not a 20% probability threshold."),
        paragraph("The cost-sensitivity table tests missed-bankruptcy weights of 5, 10 and 20 false-positive units. These illustrative weights are not currency estimates. "
                  "They expose the sensitivity of a decision rule to business priorities without replacing the capacity-based primary policy."),
    ]),
    ("Financial drivers and model sensitivity", [
        paragraph(driver_text),
        table(["Ratio", "Odds ratio per IQR", "Positive in CV fits"], [
            [x['description'], num(x['calibrated_odds_ratio_per_iqr'], 2), pct(x['cv_positive_fraction'], 0)] for x in drivers[:6]
        ]),
        paragraph("Odds ratios use an interquartile increase in the transformed development ratio and include the final calibration slope. "
                  "Ridge shrinkage addresses numerical instability from correlated inputs, but does not make each coefficient an independent economic mechanism. "
                  "The full table includes weak or counterintuitive conditional associations rather than suppressing them."),
        table(["Half IQR ratio increase", "Logistic change pp", "Forest change pp"], [
            [rows[0]['description'], num(float(rows[0]['mean_risk_change']) * 100, 2), num(float(rows[1]['mean_risk_change']) * 100, 2)]
            for rows in [[x for x in sensitivity if x['feature'] == v and float(x['iqr_shift']) > 0] for v in ['Attr1', 'Attr2', 'Attr4', 'Attr21']]
        ]),
        paragraph("The table reports percentage-point changes in average holdout scores after increasing one ratio by half its development interquartile range. "
                  "Logistic regression models log odds as linear in the transformed inputs, while the forest can capture nonlinear relationships. "
                  "The supporting sensitivity table and figure also document the model response; decreases are included in the CSV."),
        paragraph("These one-at-a-time perturbations are bounded by development support. They do not maintain all accounting identities and cannot be interpreted "
                  "as a balance-sheet stress test or the causal impact of changing a firm's finances. A coherent stress framework would require underlying statements and scenario assumptions."),
    ]),
    ("Horizon comparison and practical use", [
        table(["Horizon", "Logistic AUC", "Forest AUC", "Logistic recall", "Forest recall"], [
            [f"{h} " + ("year" if h == 1 else "years"), num([x for x in perf if x['horizon_years'] == str(h)][0]['roc_auc']),
             num([x for x in perf if x['horizon_years'] == str(h)][1]['roc_auc']),
             pct([x for x in perf if x['horizon_years'] == str(h)][0]['recall']),
             pct([x for x in perf if x['horizon_years'] == str(h)][1]['recall'])] for h in range(1, 6)
        ]),
        paragraph("The strongest discrimination is in the one-year sample. Longer-horizon performance is weaker and varies across files. "
                  "Because their company coverage and event mix differ, these results do not isolate the causal effect of time to bankruptcy. "
                  "Recall is evaluated at a separate development-selected threshold for each file and model."),
        paragraph(limitations),
        paragraph("I would use the scores to prioritize financial review, then examine liquidity, refinancing needs, earnings quality, collateral and recent developments. "
                  "A current lending application would first require a dated external cohort, firm-level identifiers, review-cost information and monitoring of calibration and data drift. "
                  "Further research should test temporal transportability and coherent statement-based stress scenarios rather than add complexity solely to raise a benchmark score."),
    ]),
    ("Execution and supporting materials", [
        paragraph("The project contains the five UCI datasets, the executable R entry point, shared functions, data diagnostics, "
                  "model validation, figures, output tables, methodology notes, this report and a one-page executive summary. "
                  "The self-contained HTML report embeds its figures and can be opened directly in a browser."),
        paragraph("From the project directory, run Rscript final_project.R, then Rscript tests/verify_results.R. "
                  "Run python3 scripts/build_reports.py to regenerate the Markdown, HTML and editable Word documents from the executed tables. "
                  "The --pdf option uses LibreOffice to regenerate the PDFs. Exact software versions are recorded in outputs/tables/software-versions.csv and docs/session-info.txt."),
        paragraph("The included datasets and relative file paths make the project self-contained. The R pipeline creates its output folders and execution cache when needed. "
                  "Screening conclusions use independently evaluated holdout predictions. No financial losses, market returns or real credit decisions are inferred from the bankruptcy labels."),
        paragraph("Sources"),
        paragraph("Tomczak, S. (2016). Polish Companies Bankruptcy. UCI Machine Learning Repository. https://doi.org/10.24432/C5F600. Data licensed under CC BY 4.0."),
        paragraph("UCI data description and financial definitions: https://archive.ics.uci.edu/dataset/365/polish+companies+bankruptcy+data"),
        paragraph("glmnet documentation: https://glmnet.stanford.edu/articles/glmnet.html"),
        paragraph("randomForest documentation: https://search.r-project.org/CRAN/refmans/randomForest/html/randomForest.html"),
    ]),
]


def setup_doc(title, subtitle=None):
    doc = Document()
    doc.core_properties.author = AUTHOR
    doc.core_properties.last_modified_by = AUTHOR
    doc.core_properties.title = title
    doc.core_properties.subject = "Bankruptcy risk analysis and model validation"
    doc.core_properties.comments = ""
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.75); section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        doc.styles[name].font.name = 'Calibri'
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(4)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
    doc.styles['Caption'].font.color.rgb = RGBColor(70, 70, 70)
    doc.styles['Caption'].font.bold = False
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(AUTHOR, 'Subtitle')
    if subtitle:
        doc.add_paragraph(subtitle)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(); run.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    run._r.append(field)
    return doc


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    n = len(headers)
    widths = [3.5] + [(7 - 3.5) / (n - 1)] * (n - 1) if n <= 3 else [2.5] + [(7 - 2.5) / (n - 1)] * (n - 1)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    border = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9')
        border.append(el)
    t._tbl.tblPr.append(border)
    for r_idx, values in enumerate([headers] + rows):
        cells = t.rows[0].cells if r_idx == 0 else t.add_row().cells
        if r_idx == 0:
            repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for c_idx, (cell, text) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[c_idx]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(str(text)); run.font.size = Pt(10)
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), '80'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '163A5F' if r_idx == 0 else ('F2F5F8' if r_idx % 2 == 0 else 'FFFFFF'))
            tcpr.append(shading)
            if r_idx == 0:
                run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
            no_split = OxmlElement('w:cantSplit'); cell._tc.getparent().get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_blocks(doc, blocks):
    for item in blocks:
        if item[0] == 'p':
            doc.add_paragraph(item[1])
        elif item[0] == 'table':
            add_table(doc, item[1], item[2])
        else:
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            width = 6.0 if item[1] == 'financial-ratio-model-sensitivity' else 6.9
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT / 'outputs/figures' / (item[1] + '.png')), width=Inches(width))
            caption = doc.add_paragraph(item[2])
            caption.style = doc.styles['Caption']; caption.runs[0].font.size = Pt(9)


def markdown_blocks(blocks, prefix='../'):
    out = []
    for item in blocks:
        if item[0] == 'p':
            out.append(item[1])
        elif item[0] == 'table':
            out.append('| ' + ' | '.join(item[1]) + ' |\n| ' + ' | '.join(['---'] * len(item[1])) + ' |\n' +
                       '\n'.join('| ' + ' | '.join(str(x) for x in row) + ' |' for row in item[2]))
        else:
            out.append(f"![{item[2]}]({prefix}outputs/figures/{item[1]}.png)\n\n{item[2]}")
    return '\n\n'.join(out)


def html_blocks(blocks):
    out = []
    for item in blocks:
        if item[0] == 'p':
            out.append('<p>' + html.escape(item[1]) + '</p>')
        elif item[0] == 'table':
            out.append('<div class="table-wrap"><table><thead><tr>' + ''.join('<th>' + html.escape(x) + '</th>' for x in item[1]) + '</tr></thead><tbody>')
            for row in item[2]:
                out.append('<tr>' + ''.join('<td>' + html.escape(str(x)) + '</td>' for x in row) + '</tr>')
            out.append('</tbody></table></div>')
        else:
            b64 = base64.b64encode((ROOT / 'outputs/figures' / (item[1] + '.png')).read_bytes()).decode()
            out.append(f'<figure><img src="data:image/png;base64,{b64}" alt="{html.escape(item[2])}"><figcaption>{html.escape(item[2])}</figcaption></figure>')
    return '\n'.join(out)


def build():
    report = setup_doc(TITLE, 'Financial ratio analysis and out of sample validation')
    for i, (title, blocks) in enumerate(sections):
        if i:
            report.add_page_break()
        report.add_heading(title, level=1)
        add_blocks(report, blocks)
    report_path = ROOT / 'Final_Project_Report_Florien_Siakoua_Toukam.docx'
    report.save(report_path)
    summary = setup_doc('Bankruptcy Risk Executive Summary')
    summary.styles['Normal'].font.size = Pt(10.5)
    summary.styles['Normal'].paragraph_format.space_after = Pt(6)
    summary_blocks = [
        paragraph(f"Objective and scope. Prioritize financial review using {integer(inventory[4]['observations'])} one-year-ahead records, "
                  f"including {integer(inventory[4]['bankruptcies'])} bankruptcies. Development and holdout are separated by exact predictor groups; "
                  "model fitting and calibration use nested, repeated cross-validation within development."),
        table(['Holdout metric', 'Ridge logistic', 'Balanced forest'], metric_rows[:5]),
        paragraph(f"Screening. At thresholds of {pct(lr['threshold'], 2)} and {pct(rf['threshold'], 2)}, "
                  f"logistic regression flags {integer(number(lr, 'tp') + number(lr, 'fp'))} records, misses {integer(lr['fn'])} bankruptcies and generates {integer(lr['fp'])} false positives. "
                  f"The forest flags {integer(number(rf, 'tp') + number(rf, 'fp'))}, misses {integer(rf['fn'])} and generates {integer(rf['fp'])} false positives. "
                  "Thresholds implement a development review budget of 20%, not automatic credit decisions."),
        paragraph(f"Risk segmentation. Low / Moderate / High bands contain {integer(bands[0]['n'])} / {integer(bands[1]['n'])} / {integer(bands[2]['n'])} records. "
                  f"Their observed bankruptcy rates are {pct(bands[0]['observed_bankruptcy_rate'])} / {pct(bands[1]['observed_bankruptcy_rate'])} / {pct(bands[2]['observed_bankruptcy_rate'])}."),
        paragraph("Financial interpretation. Lower sales growth, weaker working capital, smaller asset scale and lower profitability are associated with higher estimated odds; "
                  "higher liabilities relative to assets also increase estimated odds, conditional on the other inputs."),
        paragraph(f"Model judgment. Keep logistic regression as the interpretable reference and the forest as challenger. "
                  f"The forest's AUC advantage has a 95% interval of {num(delta['lower_95'])} to {num(delta['upper_95'])}; "
                  f"logistic regression has higher average precision and lower calibrated Brier error. "
                  f"Mean cross-validation AUCs are {num(cv_means[lr['model']])} / {num(cv_means[rf['model']])}."),
        paragraph("Limitations. Historical selected data and missing firm IDs, observation dates and loss information limit transfer to current credit decisions. "
                  "Exact duplicate grouping cannot identify repeated firms with different statements. Risk scores and bands are relative to this dataset; "
                  "current use would require dated external validation and measured review costs."),
    ]
    add_blocks(summary, summary_blocks)
    summary_path = ROOT / 'reports/Bankruptcy_Risk_Executive_Summary.docx'
    summary.save(summary_path)
    (ROOT / 'reports/underwriting-summary.md').write_text('# Bankruptcy Risk Executive Summary\n\n' + AUTHOR + '\n\n' + markdown_blocks(summary_blocks) + '\n')
    md = '# ' + TITLE + '\n\n' + AUTHOR + '\n\n' + '\n\n'.join('## ' + title + '\n\n' + markdown_blocks(blocks) for title, blocks in sections)
    (ROOT / 'reports/analysis-report.md').write_text(md + '\n')
    nav = ''.join(f'<a href="#s{i}">{html.escape(title)}</a>' for i, (title, _) in enumerate(sections))
    body = ''.join(f'<section id="s{i}"><h2>{html.escape(title)}</h2>{html_blocks(blocks)}</section>' for i, (title, blocks) in enumerate(sections))
    html_doc = '''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Polish Companies Bankruptcy Analysis</title><style>
*{box-sizing:border-box}body{margin:0;background:#f5f7f9;color:#172b3d;font:16px/1.65 system-ui,sans-serif}
header{padding:40px max(24px,calc((100vw - 1120px)/2));background:#163a5f;color:white}h1{font-size:32px;line-height:1.25;margin:0 0 12px}header p{margin:0}
.layout{max-width:1160px;margin:32px auto;display:grid;grid-template-columns:220px 1fr;gap:30px;padding:0 20px}nav{position:sticky;top:24px;align-self:start}
nav a{display:block;padding:7px 0;color:#163a5f;font-size:14px;text-decoration:none}nav a:hover{text-decoration:underline}section{background:white;padding:24px 30px;margin:0 0 24px;border:1px solid #dce4ea}
h2{font-size:24px;line-height:1.3;margin:0 0 18px}p{margin:0 0 16px}.table-wrap{overflow:auto;margin:22px 0}table{width:100%;border-collapse:collapse;font-size:14px}
td,th{border:1px solid #d9d9d9;padding:9px 12px;text-align:right}td:first-child,th:first-child{text-align:left}th{background:#163a5f;color:white}tr:nth-child(even){background:#f2f5f8}
figure{margin:24px 0}img{width:100%;height:auto}figcaption{font-size:13px;color:#526477}footer{padding:24px;text-align:center;font-size:13px}
@media(max-width:800px){.layout{display:block}nav{position:static;margin-bottom:20px}section{padding:20px}h1{font-size:26px}}
@media print{nav{display:none}.layout{display:block;margin:0}section{break-inside:avoid;border:0}header{background:white;color:black}body{background:white}}
</style></head><body>'''
    html_doc += f'<header><h1>{TITLE}</h1><p>{AUTHOR} · Financial ratio analysis and out of sample validation</p></header><div class="layout"><nav aria-label="Contents">{nav}</nav><main>{body}</main></div><footer>Data: Tomczak (2016), UCI Machine Learning Repository · DOI 10.24432/C5F600</footer></body></html>'
    (ROOT / 'reports/analysis-report.html').write_text(html_doc)

    readme = f'''# Polish Companies Bankruptcy Analysis

Financial ratio analysis for bankruptcy screening, with repeated cross-validation, a separate holdout, risk segmentation and a calibrated random-forest challenger.

{headline}

## One year holdout results

{markdown_blocks([table(['Metric', 'Ridge logistic', 'Balanced forest'], metric_rows[:6])], '')}

The holdout contains **{integer(lr['n'])} records and {integer(lr['bankruptcies'])} bankruptcies**. Mean development cross-validation AUCs are {num(cv_means[lr['model']])} and {num(cv_means[rf['model']])}, respectively. The paired forest-minus-logistic holdout AUC difference is {num(delta['estimate'])} (95% interval {num(delta['lower_95'])} to {num(delta['upper_95'])}).

![Holdout ROC and precision recall curves](outputs/figures/holdout-roc-and-precision-recall.png)

## Screening and risk bands

{policy_text}

{markdown_blocks([table(['Logistic band', 'Records', 'Mean predicted risk', 'Observed bankruptcy rate'], [[x['band'], integer(x['n']), pct(x['average_predicted_risk']), pct(x['observed_bankruptcy_rate'])] for x in bands])], '')}

Band boundaries are {pct(edges[0]['threshold'], 2)} and {pct(edges[1]['threshold'], 2)}, determined from development scores. Actual holdout review shares are {pct(lr['review_rate'])} and {pct(rf['review_rate'])}; a fixed threshold does not force an exact holdout review percentage.

![Holdout risk bands](outputs/figures/holdout-bankruptcy-risk-bands.png)

{comparison}

## Financial interpretation

{driver_text}

{markdown_blocks([table(['Ratio', 'Calibrated odds ratio per IQR'], [[x['description'], num(x['calibrated_odds_ratio_per_iqr'], 2)] for x in drivers[:6]])], '')}

The [full driver table](outputs/tables/logistic-risk-drivers.csv) reports all 16 ratios, including small and less intuitive associations. [Ratio sensitivity](outputs/tables/financial-ratio-sensitivity.csv) measures isolated changes within development support; it is not an accounting-consistent stress test.

## Data quality and sample inclusion

The [UCI Polish Companies Bankruptcy data](https://doi.org/10.24432/C5F600) contain five separate forecasting samples. **5year.arff is the one-year-ahead sample; 1year.arff is the five-year-ahead sample.** The five source files are included unchanged under CC BY 4.0 with [attribution and checksums](docs/data-source.md). Each has 64 predictors plus the bankruptcy indicator.

{selection_text}

![Development sample exclusion diagnostic](outputs/figures/complete-case-selection-impact.png)

Bankruptcy prevalence is {pct(inventory[4]['bankruptcy_rate'])} in the one-year sample, making accuracy alone insufficient. [Sample-selection diagnostics](outputs/tables/complete-case-selection-diagnostic.csv) cover all five horizons. The [development correlation matrix](outputs/figures/development-ratio-correlations.png) and [VIF table](outputs/tables/development-collinearity-diagnostics.csv) describe dependence among the model ratios without selecting features from holdout outcomes.

## Validation design

1. Group identical raw predictor rows and reserve approximately 20% of groups for holdout evaluation.
2. Use two repeats of grouped five-fold cross-validation in development, with a nested three-fold calibration step.
3. Estimate median imputation, missing indicators and 1st/99th percentile clipping inside each training sample.
4. Fit ridge logistic regression and a balanced 500-tree random forest using the same 16 finance-defined ratios.
5. Set separate calibrated screening thresholds from a 20% development review budget; define logistic risk bands from development scores.
6. Evaluate the fixed rules on holdout, including ROC AUC, average precision, calibration, confusion counts, lift and paired group-bootstrap uncertainty.

The fixed ridge penalty is 0.01; no hyperparameter or threshold selection uses holdout outcomes. The five horizons are analyzed separately, not pooled. Full details are in [methodology notes](docs/methodology-notes.md).

## Results and reports

- [One-page executive summary](reports/underwriting-summary.md)
- [Complete analysis report](reports/analysis-report.md)
- [Complete PDF report](Final_Project_Report_Florien_Siakoua_Toukam.pdf)
- [Downloadable self-contained HTML report](reports/analysis-report.html) — download and open locally; GitHub does not render HTML reports inline.
- [Holdout results by horizon](outputs/tables/horizon-holdout-performance.csv)
- [Cross-validation fold results](outputs/tables/cross-validation-fold-results.csv)
- [Development capacity analysis](outputs/tables/development-capacity-analysis.csv)
- [Illustrative cost sensitivity](outputs/tables/development-illustrative-cost-sensitivity.csv)
- [Holdout confidence intervals](outputs/tables/holdout-metric-confidence-intervals.csv)
- [All figures](outputs/figures/) and [all tables](outputs/tables/)

## Limitations and further work

{limitations}

The 16-ratio specification and fixed regularization are intentionally limited; broader feature or tuning searches have not been evaluated. Bootstrap intervals are conditional on the fitted models and split. Risk bands are cohort-relative, and cost weights are illustrative. Multi-horizon differences also reflect differing samples.

The next substantive steps are dated external validation, firm-level grouping, measured review and loss costs, drift monitoring and accounting-consistent financial scenarios. Further model complexity should be justified by those tests.

## Run the analysis

Use R 4.5.2 and open the RStudio project or start a terminal in this folder. Restore the package versions recorded in [renv.lock](renv.lock):

```r
install.packages("renv", repos = "https://cloud.r-project.org")
renv::restore(lockfile = "renv.lock", library = ".cache/R-library", prompt = FALSE)
.libPaths(c(".cache/R-library", .libPaths()))
```

```sh
R_LIBS_USER=.cache/R-library Rscript final_project.R
R_LIBS_USER=.cache/R-library Rscript tests/verify_results.R
```

The terminal commands above use macOS/Linux syntax. In RStudio or on Windows, run the R setup above, then `source("final_project.R")` and `source("tests/verify_results.R")`. The lockfile records R and package versions; it does not install R or system compilers. If an archived package has no compatible binary, source installation requires the platform's R build tools, including a Fortran compiler for glmnet. The local package library stays outside the published ZIP through the existing `.cache/` exclusion.

This regenerates all analytical CSVs and PNGs, including their folders if absent. All required data are included, and no external local files or saved model cache are required. The tested R and package versions are recorded in [software versions](outputs/tables/software-versions.csv) and [session information](docs/session-info.txt). Numerical fitting checks stop the analysis if convergence fails. Both macOS Quartz and Cairo-capable PNG rendering are supported.

To rebuild the written reports after a successful analysis:

```sh
python3 -m pip install -r requirements.txt
python3 scripts/build_reports.py
```

To regenerate PDFs as well, install LibreOffice and make its `soffice` command available, then run `python3 scripts/build_reports.py --pdf`. Otherwise open the generated Word files and export them to PDF. Run `python3 tests/verify_package.py` after rebuilding. The HTML report embeds its figures and does not need an internet connection to display the analysis.

## Structure

```text
final_project.R
polish-companies-bankruptcy-analysis.Rproj
R/
  00_functions.R
  01_data_preparation_and_diagnostics.R
  02_model_validation.R
  03_visual_summaries.R
polish+companies+bankruptcy+data/     # Five unchanged ARFF files
docs/                              # Data source, checksums, methods, R session
outputs/figures/                    # PNG visual summaries
outputs/tables/                     # Results, predictions and diagnostics
reports/                           # Markdown, HTML and executive summary
scripts/build_reports.py
tests/verify_results.R
tests/verify_package.py
requirements.txt
renv.lock
Final_Project_Report_Florien_Siakoua_Toukam.docx
Final_Project_Report_Florien_Siakoua_Toukam.pdf
```

## Sources

Tomczak, S. (2016). *Polish Companies Bankruptcy*. UCI Machine Learning Repository. [DOI 10.24432/C5F600](https://doi.org/10.24432/C5F600). Data: [CC BY 4.0](https://creativecommons.org/licenses/by/4.0/).

Methods and implementation: [glmnet](https://glmnet.stanford.edu/articles/glmnet.html), [randomForest](https://search.r-project.org/CRAN/refmans/randomForest/html/randomForest.html), and [pROC](https://xrobin.github.io/pROC/).
'''
    (ROOT / 'README.md').write_text(readme)
    return [report_path, summary_path]


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--pdf', action='store_true')
    args = parser.parse_args()
    paths = build()
    if args.pdf:
        executable = shutil.which('soffice')
        if not executable:
            raise SystemExit('Word, Markdown and HTML reports created. PDF conversion requires soffice on PATH.')
        with tempfile.TemporaryDirectory(prefix='bankruptcy-report-') as tmp:
            for path in paths:
                subprocess.run([executable, '-env:UserInstallation=' + Path(tmp).as_uri(), '--headless',
                                '--convert-to', 'pdf', '--outdir', str(path.parent), str(path)], check=True)
                if not path.with_suffix('.pdf').is_file():
                    raise RuntimeError('PDF conversion did not produce ' + str(path.with_suffix('.pdf')))
    print('Written reports rebuilt from the executed tables.')
